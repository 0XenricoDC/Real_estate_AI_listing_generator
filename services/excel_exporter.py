"""
Servizio per l'esportazione dei lead in formato Excel
"""

from typing import List
from datetime import datetime
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.lead import Lead
from models.search_filters import SearchFilters
from config import EXCEL_COLORS, PRIORITY_THRESHOLDS


class ExcelExporter:
    """
    Servizio per esportare i lead in un file Excel formattato.
    Crea un workbook con fogli separati per priorità e statistiche.
    """

    # Intestazioni colonne
    HEADERS = [
        'Score', 'Priorità', 'Nome', 'Cognome', 'Telefono', 'Email',
        'Indirizzo', 'Comune', 'Provincia', 'Prezzo', 'Tipo',
        'MQ', 'Locali', 'Stato', 'Data Annuncio', 'Fonte', 'URL', 'Note'
    ]

    # Larghezze colonne
    COLUMN_WIDTHS = {
        'A': 8,   # Score
        'B': 10,  # Priorità
        'C': 15,  # Nome
        'D': 15,  # Cognome
        'E': 15,  # Telefono
        'F': 25,  # Email
        'G': 30,  # Indirizzo
        'H': 15,  # Comune
        'I': 10,  # Provincia
        'J': 15,  # Prezzo
        'K': 15,  # Tipo
        'L': 8,   # MQ
        'M': 8,   # Locali
        'N': 15,  # Stato
        'O': 12,  # Data
        'P': 12,  # Fonte
        'Q': 40,  # URL
        'R': 30   # Note
    }

    def __init__(self):
        self.workbook = None

    def export(self, leads: List[Lead], filters: SearchFilters, output_path: str = None) -> str:
        """
        Esporta i lead in un file Excel.

        Args:
            leads: Lista di lead da esportare
            filters: Filtri utilizzati per la ricerca
            output_path: Percorso del file di output (opzionale)

        Returns:
            Percorso del file Excel creato
        """
        if not output_path:
            # Genera nome file con data e località
            date_str = datetime.now().strftime('%Y%m%d_%H%M')
            localita_clean = filters.localita.replace(' ', '_').lower()
            output_path = f"acquisizioni_{date_str}_{localita_clean}.xlsx"

        self.workbook = Workbook()

        # Rimuovi il foglio di default
        default_sheet = self.workbook.active
        self.workbook.remove(default_sheet)

        # Separa i lead per priorità
        alta = [l for l in leads if l.priorita == 'ALTA']
        media = [l for l in leads if l.priorita == 'MEDIA']
        bassa = [l for l in leads if l.priorita == 'BASSA']

        # Crea i fogli
        self._create_leads_sheet('Lead Alta Priorità', alta, 'alta')
        self._create_leads_sheet('Lead Media Priorità', media, 'media')
        self._create_leads_sheet('Lead Bassa Priorità', bassa, 'bassa')
        self._create_leads_sheet('Tutti i Lead', leads, None)
        self._create_statistics_sheet(leads, filters)

        # Salva il file
        self.workbook.save(output_path)

        return output_path

    def _create_leads_sheet(self, sheet_name: str, leads: List[Lead], priority_type: str):
        """Crea un foglio con la lista dei lead"""
        ws = self.workbook.create_sheet(sheet_name)

        # Stili
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color=EXCEL_COLORS['header'],
                                  end_color=EXCEL_COLORS['header'],
                                  fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Intestazioni
        for col, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Dati
        for row_num, lead in enumerate(leads, 2):
            row_data = lead.to_excel_row()

            # Colore della riga basato sulla priorità
            if lead.priorita == 'ALTA':
                row_fill = PatternFill(start_color=EXCEL_COLORS['alta'],
                                       end_color=EXCEL_COLORS['alta'],
                                       fill_type='solid')
            elif lead.priorita == 'MEDIA':
                row_fill = PatternFill(start_color=EXCEL_COLORS['media'],
                                       end_color=EXCEL_COLORS['media'],
                                       fill_type='solid')
            else:
                row_fill = PatternFill(start_color=EXCEL_COLORS['bassa'],
                                       end_color=EXCEL_COLORS['bassa'],
                                       fill_type='solid')

            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.fill = row_fill
                cell.border = thin_border

                # Allineamento specifico per alcune colonne
                if col in [1, 2, 12, 13]:  # Score, Priorità, MQ, Locali
                    cell.alignment = Alignment(horizontal='center')

        # Imposta larghezza colonne
        for col_letter, width in self.COLUMN_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        # Blocca la prima riga
        ws.freeze_panes = 'A2'

        # Aggiungi filtri automatici
        if leads:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(self.HEADERS))}{len(leads) + 1}"

    def _create_statistics_sheet(self, leads: List[Lead], filters: SearchFilters):
        """Crea il foglio delle statistiche"""
        ws = self.workbook.create_sheet('Statistiche')

        # Stili
        title_font = Font(bold=True, size=14)
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color=EXCEL_COLORS['header'],
                                  end_color=EXCEL_COLORS['header'],
                                  fill_type='solid')

        # Titolo
        ws['A1'] = 'RIEPILOGO RICERCA'
        ws['A1'].font = title_font
        ws.merge_cells('A1:D1')

        # Data ricerca
        ws['A3'] = 'Data ricerca:'
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M')
        ws['A3'].font = header_font

        # Filtri utilizzati
        ws['A5'] = 'FILTRI UTILIZZATI'
        ws['A5'].font = header_font
        ws.merge_cells('A5:D5')

        ws['A6'] = 'Località:'
        ws['B6'] = filters.localita

        ws['A7'] = 'Raggio:'
        ws['B7'] = f"{filters.raggio_km} km"

        ws['A8'] = 'Prezzo:'
        ws['B8'] = filters.get_prezzo_display()

        ws['A9'] = 'Tipo immobile:'
        ws['B9'] = filters.tipo_immobile

        ws['A10'] = 'Siti cercati:'
        ws['B10'] = ', '.join(filters.siti_attivi)

        # Statistiche lead
        ws['A12'] = 'STATISTICHE LEAD'
        ws['A12'].font = header_font
        ws.merge_cells('A12:D12')

        alta = sum(1 for l in leads if l.priorita == 'ALTA')
        media = sum(1 for l in leads if l.priorita == 'MEDIA')
        bassa = sum(1 for l in leads if l.priorita == 'BASSA')
        scores = [l.lead_score for l in leads] if leads else [0]
        prezzi = [l.prezzo_richiesto for l in leads if l.prezzo_richiesto]

        stats = [
            ('Totale lead trovati:', len(leads)),
            ('Lead alta priorità:', alta),
            ('Lead media priorità:', media),
            ('Lead bassa priorità:', bassa),
            ('Score medio:', f"{sum(scores) / len(scores):.1f}" if scores else "N/A"),
            ('Score massimo:', max(scores) if scores else "N/A"),
            ('Score minimo:', min(scores) if scores else "N/A"),
            ('Prezzo medio:', f"€ {sum(prezzi) / len(prezzi):,.0f}".replace(',', '.') if prezzi else "N/A"),
        ]

        for i, (label, value) in enumerate(stats, 13):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = header_font

        # Distribuzione per fonte
        ws['A22'] = 'DISTRIBUZIONE PER FONTE'
        ws['A22'].font = header_font
        ws.merge_cells('A22:D22')

        fonti = {}
        for lead in leads:
            fonte = lead.fonte or 'Sconosciuta'
            fonti[fonte] = fonti.get(fonte, 0) + 1

        for i, (fonte, count) in enumerate(sorted(fonti.items()), 23):
            ws[f'A{i}'] = fonte
            ws[f'B{i}'] = count

        # Larghezza colonne
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15

    def export_simple(self, leads: List[Lead], output_path: str) -> str:
        """
        Esporta i lead in un file Excel semplice (singolo foglio).

        Args:
            leads: Lista di lead da esportare
            output_path: Percorso del file di output

        Returns:
            Percorso del file Excel creato
        """
        self.workbook = Workbook()
        ws = self.workbook.active
        ws.title = 'Lead'

        # Intestazioni
        header_font = Font(bold=True)
        for col, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font

        # Dati
        for row_num, lead in enumerate(leads, 2):
            row_data = lead.to_excel_row()
            for col, value in enumerate(row_data, 1):
                ws.cell(row=row_num, column=col, value=value)

        # Larghezza colonne
        for col_letter, width in self.COLUMN_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        self.workbook.save(output_path)
        return output_path
